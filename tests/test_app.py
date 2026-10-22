# tests/test_app.py — Pruebas automáticas de Academia Horizonte
# Estas pruebas usan el cliente de Flask: simulan un navegador sin abrir la app.
# No usan el archivo .env: las credenciales se definen aquí antes de importar app.

import io
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# 1. Configuración de prueba, ANTES de importar la aplicación.
#    app.py lee estas variables al cargarse; por eso se definen primero.
# ---------------------------------------------------------------------------
ADMIN_CORREO = "admin@prueba.local"
ADMIN_CLAVE = "ClaveAdminPrueba123"
USUARIO_CORREO = "usuario@prueba.local"
USUARIO_CLAVE = "ClaveUsuarioPrueba123"
SECRET_KEY = "clave-secreta-solo-para-pruebas"

os.environ["ADMIN_CORREO"] = ADMIN_CORREO
os.environ["ADMIN_CLAVE"] = ADMIN_CLAVE
os.environ["USUARIO_CORREO"] = USUARIO_CORREO
os.environ["USUARIO_CLAVE"] = USUARIO_CLAVE
os.environ["SECRET_KEY"] = SECRET_KEY

# Se quita una ruta externa para que la prueba no dependa de la máquina
os.environ.pop("CARPETA_INSUMOS", None)

# ---------------------------------------------------------------------------
# 2. Ahora sí se importa la app, ya con las credenciales de prueba puestas.
# ---------------------------------------------------------------------------
import app as aplicacion  # noqa: E402

# Ruta de la raíz del proyecto (la carpeta que contiene app.py)
RUTA_RAIZ = Path(__file__).resolve().parents[1]


@pytest.fixture()
def cliente():
    """Entrega un navegador falso con su propia sesión y cookies."""
    aplicacion.app.config["TESTING"] = True
    with aplicacion.app.test_client() as navegador:
        yield navegador


def entrar(navegador, correo, clave):
    """Simula escribir el correo y la contraseña en la pantalla de entrada."""
    return navegador.post("/login", data={"correo": correo, "clave": clave})


def test_sin_sesion_redirige_a_login(cliente):
    """Simula abrir / sin haber entrado: debe mandar a la pantalla de entrada."""
    respuesta = cliente.get("/")

    # 302 significa "redirigir"; la cabecera Location dice hacia dónde
    assert respuesta.status_code == 302
    assert respuesta.headers["Location"].endswith("/login")


def test_clave_incorrecta_no_entra(cliente):
    """Simula escribir un correo y una clave equivocados en /login."""
    respuesta = cliente.post(
        "/login",
        data={"correo": ADMIN_CORREO, "clave": "esta-clave-no-es"},
    )

    # La pantalla se vuelve a mostrar (200) con el mensaje de error
    assert respuesta.status_code == 200
    html = respuesta.get_data(as_text=True)
    assert "Correo o contraseña incorrectos." in html

    # Además, la sesión sigue vacía: / vuelve a redirigir a /login
    assert cliente.get("/").status_code == 302


def test_usuario_no_puede_cargar(cliente):
    """Simula entrar como usuario normal e intentar subir Excel a /cargar."""
    respuesta_entrada = entrar(cliente, USUARIO_CORREO, USUARIO_CLAVE)
    assert respuesta_entrada.status_code == 302  # entró bien

    # /cargar es solo para admin: debe responder 403 (prohibido)
    respuesta = cliente.post("/cargar", data={"accion": "ver"})
    assert respuesta.status_code == 403


def test_excel_modificado_cambia_contadores(cliente):
    """Simula subir un registro con las notas de Camila Rojas en 40.

    Se lee el Excel original de Insumos, se cambia solo la copia en memoria
    y se sube esa copia. El archivo original nunca se modifica.
    """
    carpeta_insumos = RUTA_RAIZ / "Insumos"
    ruta_maestro = carpeta_insumos / "Maestro_Estudiantes.xlsx"
    ruta_registro = carpeta_insumos / "Registro_Evaluaciones.xlsx"

    # Si no existen los datos reales, se avisa y se omite esta prueba
    if not ruta_maestro.exists() or not ruta_registro.exists():
        pytest.skip("Faltan los Excel de Insumos/ para esta prueba.")

    # Se leen los bytes del original y se abre una COPIA en memoria
    datos_maestro = ruta_maestro.read_bytes()
    datos_registro = ruta_registro.read_bytes()
    libro = load_workbook(io.BytesIO(datos_registro))

    # BUCLE: busca las filas de Camila (identificación 101230456) y pone Nota = 40
    filas_camila = 0
    for fila in libro.active.iter_rows(min_row=2):
        if str(fila[0].value).strip() == "101230456":
            fila[3].value = 40  # columna D: Nota
            filas_camila += 1
    assert filas_camila == 4  # la estudiante tiene 4 módulos registrados

    # El Excel modificado se guarda en otro BytesIO, no en el disco
    registro_modificado = io.BytesIO()
    libro.save(registro_modificado)
    registro_modificado.seek(0)

    # Se entra como administrador
    respuesta_entrada = entrar(cliente, ADMIN_CORREO, ADMIN_CLAVE)
    assert respuesta_entrada.status_code == 302

    # Simula el formulario del panel: los dos archivos y "Procesar y ver"
    respuesta = cliente.post(
        "/cargar",
        data={
            "accion": "ver",
            "maestro": (io.BytesIO(datos_maestro), "Maestro_Estudiantes.xlsx"),
            "registro": (registro_modificado, "Registro_Evaluaciones.xlsx"),
        },
        content_type="multipart/form-data",
    )

    assert respuesta.status_code == 200
    html = respuesta.get_data(as_text=True)

    # Camila baja de 91.25 a 40: deja de aprobar y pasa a participación.
    # Por eso los contadores cambian de 19/3/1 a 18/4/1.
    assert "Aprobación: 18" in html
    assert "Participación: 4" in html
    assert "Sin certificado: 1" in html
